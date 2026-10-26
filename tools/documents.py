"""
tools/documents.py
---------------------
Stage 9: actually generate files, not just folders — plain text,
real Word documents, PDFs, and Excel spreadsheets.

docx/pdf/xlsx support needs python-docx / fpdf2 / openpyxl respectively
(see requirements.txt). Same optional-dependency pattern as voice/AI
packages elsewhere in the project: imported lazily inside each
function, with a friendly error (not a crash) if the package isn't
installed, so a missing package only disables that one file format.

Overwriting an existing file is the only "destructive" part of this —
there's no delete here — so creating a document that would clobber an
existing file requires confirmation; creating a brand new file doesn't.
"""

import csv
import io
from pathlib import Path

from tools.files import _resolve_path  # same shortcuts (downloads/desktop/etc.) as every other file tool


def _check_overwrite(target: Path, confirmed: bool) -> dict | None:
    """Returns a confirmation_required dict if target exists and the
    caller hasn't confirmed yet, else None (safe to proceed)."""
    if target.exists() and not confirmed:
        return {
            "confirmation_required": True,
            "message": f"'{target}' already exists and would be overwritten. Ask the user to confirm first.",
        }
    return None


def create_text_document(path: str, content: str, confirmed: bool = False) -> dict:
    """Create a plain .txt file with the given content."""
    target = _resolve_path(path)
    guard = _check_overwrite(target, confirmed)
    if guard:
        return guard
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")
    except OSError as e:
        return {"error": f"Could not write file: {e}"}
    return {"status": "created", "path": str(target)}


def create_word_document(path: str, content: str, confirmed: bool = False) -> dict:
    """Create a real .docx Word document. content's lines become
    paragraphs (a blank line becomes a paragraph break)."""
    target = _resolve_path(path)
    if target.suffix.lower() != ".docx":
        target = target.with_suffix(".docx")
    guard = _check_overwrite(target, confirmed)
    if guard:
        return guard

    try:
        import docx
    except ImportError:
        return {"error": "python-docx isn't installed. Run: pip install -r requirements.txt"}

    try:
        document = docx.Document()
        for line in content.split("\n"):
            document.add_paragraph(line)
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(target))
    except OSError as e:
        return {"error": f"Could not write file: {e}"}
    return {"status": "created", "path": str(target)}


def create_pdf_document(path: str, content: str, confirmed: bool = False) -> dict:
    """Create a simple .pdf with the given text content, wrapped to the page."""
    target = _resolve_path(path)
    if target.suffix.lower() != ".pdf":
        target = target.with_suffix(".pdf")
    guard = _check_overwrite(target, confirmed)
    if guard:
        return guard

    try:
        from fpdf import FPDF
    except ImportError:
        return {"error": "fpdf2 isn't installed. Run: pip install -r requirements.txt"}

    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.set_auto_page_break(auto=True, margin=15)
        for line in content.split("\n"):
            # multi_cell wraps long lines instead of running off the page edge.
            pdf.multi_cell(0, 8, text=line)
        target.parent.mkdir(parents=True, exist_ok=True)
        pdf.output(str(target))
    except OSError as e:
        return {"error": f"Could not write file: {e}"}
    return {"status": "created", "path": str(target)}


def create_spreadsheet(path: str, csv_data: str, confirmed: bool = False) -> dict:
    """Create a real .xlsx spreadsheet from CSV-formatted data, e.g.
    'Name,Age\\nAlice,30\\nBob,25'."""
    target = _resolve_path(path)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")
    guard = _check_overwrite(target, confirmed)
    if guard:
        return guard

    try:
        import openpyxl
    except ImportError:
        return {"error": "openpyxl isn't installed. Run: pip install -r requirements.txt"}

    try:
        rows = list(csv.reader(io.StringIO(csv_data)))
    except csv.Error as e:
        return {"error": f"csv_data couldn't be parsed as CSV: {e}"}
    if not rows:
        return {"error": "csv_data was empty."}

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(target))
    except OSError as e:
        return {"error": f"Could not write file: {e}"}
    return {"status": "created", "path": str(target), "rows": len(rows)}
